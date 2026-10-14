"""تحويل الوثائق والاستمارات (Markdown) إلى ملفات Word جاهزة للطباعة بالهوية البصرية: من اليمين، Noto Naskh Arabic،
عناوين زرقاء وخط ذهبي، وجداول بترويسة ملونة، والأقواس وما بينها بالأحمر. يُستدعى: turath forms-gen"""
from __future__ import annotations
import re
from pathlib import Path

from . import sanitize

ROOT = Path(__file__).parent.parent
FORMS = ["04-oral-interview-form.md", "05-field-observation-form.md", "06-community-survey.md",
         "07-object-image-site-form.md", "08-documentation-guide.md", "09-study-template.md"]
BLUE, GOLD, RED, LIGHT = "1F4E79", "B8860B", "FF0000", "EAF1F8"


def md_to_docx(md: str, out: Path, title: str = "") -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    d = Document()
    for s in d.sections:
        s.right_margin = s.left_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
        bidi = OxmlElement("w:bidi")
        s._sectPr.append(bidi)
    st = d.styles["Normal"]
    st.font.name = "Noto Naskh Arabic"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:cs"), "Noto Naskh Arabic")

    def rtl(par):
        par._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def runs(par, text, size=12, color=None, bold=False):
        text = re.sub(r"`([^`]+)`", r"\1", text)
        parts = re.split(r"(\*\*[^*]+\*\*|\([^()]*\)|\[[^\[\]]*\])", text)
        for piece in parts:
            if not piece:
                continue
            b = bold
            if piece.startswith("**") and piece.endswith("**"):
                piece, b = piece[2:-2], True
            r = par.add_run(piece)
            r.font.size = Pt(size)
            r.font.bold = b
            rPr = r._r.get_or_add_rPr()
            rPr.append(OxmlElement("w:rtl"))
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int(size * 2)))
            rPr.append(szcs)
            if b:
                rPr.append(OxmlElement("w:bCs"))
            if piece.startswith(("(", "[")):
                r.font.color.rgb = RGBColor.from_string(RED)
            elif color:
                r.font.color.rgb = RGBColor.from_string(color)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), fill)
        tcPr.append(sh)

    def gold_rule(par):
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:color", GOLD), ("w:space", "4")):
            b.set(qn(k), v)
        bdr.append(b)
        pPr.append(bdr)

    head = d.add_paragraph()
    rtl(head)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(head, "إدارة الدراسات والبحوث التراثية الإماراتية", 11, GOLD, True)
    lines = md.replace("\r", "").split("\n")
    i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if block and not block[0].strip().startswith(("graph", "flowchart", "erDiagram", "%%")):
                p = d.add_paragraph()
                rtl(p)
                runs(p, "\n".join(block), 10)
            continue
        if l.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1]):
            row = lambda s: [c.strip() for c in s.strip().strip("|").split("|")]
            hdr = row(l)
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append(row(lines[i]))
                i += 1
            t = d.add_table(rows=1 + len(body), cols=len(hdr))
            t.style = "Table Grid"
            tblPr = t._tbl.tblPr
            tblPr.append(OxmlElement("w:bidiVisual"))
            for r_i, cells in enumerate([hdr] + body):
                for c_i in range(len(hdr)):
                    cell = t.rows[r_i].cells[c_i]
                    cell.paragraphs[0].text = ""
                    rtl(cell.paragraphs[0])
                    runs(cell.paragraphs[0], cells[c_i] if c_i < len(cells) else "", 11, BLUE if r_i == 0 else None, r_i == 0)
                    if r_i == 0:
                        shade(cell, LIGHT)
            d.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        p = None
        if m:
            lvl = len(m.group(1))
            p = d.add_paragraph()
            rtl(p)
            runs(p, m.group(2), {1: 20, 2: 16, 3: 14, 4: 13}[lvl], BLUE if lvl < 3 else "A0522D", True)
            if lvl == 1:
                gold_rule(p)
        elif re.match(r"^\s*[-*]\s+", l):
            p = d.add_paragraph()
            rtl(p)
            runs(p, "▪ " + re.sub(r"^\s*[-*]\s+", "", l).replace("[ ]", "☐").replace("[x]", "☑"))
        elif re.match(r"^\s*\d+[.)]\s+", l):
            p = d.add_paragraph()
            rtl(p)
            runs(p, l.strip())
        elif l.startswith(">"):
            p = d.add_paragraph()
            rtl(p)
            runs(p, l.lstrip("> "), 11, "5B6B7A")
        elif l.strip() and not re.match(r"^-{3,}$", l.strip()):
            p = d.add_paragraph()
            rtl(p)
            runs(p, l.strip())
        i += 1
    sanitize.scrub_docx(d, author="إدارة الدراسات والبحوث التراثية الإماراتية", title=title)
    import io
    buf = io.BytesIO()
    d.save(buf)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(sanitize.scrub_package(buf.getvalue()))
    return out


def generate() -> list[Path]:
    out = []
    for name in FORMS:
        src = ROOT / "docs" / name
        if src.exists():
            md = src.read_text(encoding="utf-8")
            title = next((l.lstrip("# ").strip() for l in md.splitlines() if l.startswith("# ")), name)
            out.append(md_to_docx(md, ROOT / "forms" / name.replace(".md", ".docx"), title))
    return out
